import openpyxl
from openpyxl import load_workbook

file_path = r"D:\next-postgres-starter\Prompts\RS.GE - WB_Items_IN.xlsx"
wb = load_workbook(file_path, data_only=True)
ws = wb.active

# Get first row (headers)
headers = []
for cell in ws[1]:
    headers.append(cell.value)

print("ALL COLUMNS IN EXCEL FILE:")
print("=" * 80)
for i, header in enumerate(headers):
    if header:
        print(f"Col {i}: {header}")
    else:
        print(f"Col {i}: [EMPTY]")

print("\n" + "=" * 80)
print("First 3 data rows (sample):")
print("=" * 80)
for row_idx in range(2, 5):
    print(f"\nRow {row_idx}:")
    for col_idx, header in enumerate(headers):
        if header:
            cell_value = ws.cell(row=row_idx, column=col_idx+1).value
            print(f"  {header}: {cell_value}")
