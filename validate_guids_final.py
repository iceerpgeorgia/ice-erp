"""
RS.GE WB_Items GUID Validation - Direct GUID comparison against database.
Fast validation using GUID columns only.
"""

import openpyxl
from openpyxl import load_workbook
import json
import os
import urllib.parse
from datetime import datetime
from dotenv import load_dotenv
import psycopg2

load_dotenv()

def parse_database_url(url):
    """Parse PostgreSQL connection URL."""
    try:
        url_decoded = urllib.parse.unquote(url)
        result = urllib.parse.urlparse(url_decoded)
        return {
            'host': result.hostname,
            'port': result.port or 5432,
            'user': result.username,
            'password': result.password,
            'database': result.path.lstrip('/')
        }
    except Exception as e:
        print(f"[ERROR] Error parsing DATABASE_URL: {e}")
        return None


def load_xlsx_and_extract_guids(file_path):
    """Load XLSX and extract GUID columns into memory-efficient format."""
    print(f"[FILE] Loading: {file_path}")
    
    if not os.path.exists(file_path):
        print(f"[ERROR] File not found")
        return None
    
    try:
        wb = load_workbook(file_path, data_only=True)
        ws = wb.active
        
        # Get headers
        headers = [cell.value for cell in ws[1]]
        
        # Find column indices
        cols = {
            'counteragent': headers.index('კონტრაგენტის GUID') if 'კონტრაგენტის GUID' in headers else None,
            'project': headers.index('პროექტი GUID') if 'პროექტი GUID' in headers else None,
            'inventory': headers.index('საქონელი GUID') if 'საქონელი GUID' in headers else None,
            'code': headers.index('კოდის GUID') if 'კოდის GUID' in headers else None,
        }
        
        print(f"   GUID columns: {cols}")
        
        # Extract all GUIDs
        guids_by_type = {
            'counteragent': set(),
            'project': set(),
            'inventory': set(),
            'code': set()
        }
        
        total_rows = 0
        for row_idx in range(2, ws.max_row + 1):
            total_rows += 1
            
            if total_rows % 5000 == 0:
                print(f"   Processing row {total_rows}...")
            
            for guid_type, col_idx in cols.items():
                if col_idx is not None:
                    cell_value = ws.cell(row=row_idx, column=col_idx + 1).value
                    if cell_value and str(cell_value).strip():
                        guids_by_type[guid_type].add(str(cell_value).strip().lower())
        
        print(f"   Total rows: {total_rows}")
        for guid_type, guids in guids_by_type.items():
            print(f"   - {guid_type}: {len(guids)} unique GUIDs")
        
        return guids_by_type
    
    except Exception as e:
        print(f"[ERROR] Error loading XLSX: {e}")
        import traceback
        traceback.print_exc()
        return None


def fetch_db_guids():
    """Fetch all GUIDs from database into sets."""
    print(f"\n[DB] Fetching GUIDs from database...")
    
    try:
        db_config = parse_database_url(os.getenv('DATABASE_URL'))
        if not db_config:
            return None
        
        conn = psycopg2.connect(**db_config)
        cur = conn.cursor()
        
        # Get counteragent GUIDs
        cur.execute('SELECT LOWER(counteragent_uuid::text) FROM counteragents WHERE is_active = TRUE')
        counteragent_guids = set(row[0] for row in cur.fetchall())
        print(f"   Counteragents: {len(counteragent_guids):,}")
        
        # Get project GUIDs
        cur.execute('SELECT LOWER(project_uuid::text) FROM projects')
        project_guids = set(row[0] for row in cur.fetchall())
        print(f"   Projects: {len(project_guids):,}")
        
        # Get inventory GUIDs
        cur.execute('SELECT LOWER(uuid::text) FROM inventories WHERE is_active = TRUE')
        inventory_guids = set(row[0] for row in cur.fetchall())
        print(f"   Inventories: {len(inventory_guids):,}")
        
        # Get financial code GUIDs
        cur.execute('SELECT LOWER(uuid::text) FROM financial_codes WHERE is_active = TRUE')
        code_guids = set(row[0] for row in cur.fetchall())
        print(f"   Financial codes: {len(code_guids):,}")
        
        cur.close()
        conn.close()
        
        return {
            'counteragent': counteragent_guids,
            'project': project_guids,
            'inventory': inventory_guids,
            'code': code_guids
        }
    
    except Exception as e:
        print(f"[ERROR] Error fetching GUIDs: {e}")
        import traceback
        traceback.print_exc()
        return None


def validate_guids(xlsx_guids, db_guids):
    """Validate XLSX GUIDs against database GUIDs."""
    print(f"\n[VALIDATION] Comparing GUIDs...")
    
    results = {}
    for guid_type in ['counteragent', 'project', 'inventory', 'code']:
        xlsx_set = xlsx_guids[guid_type]
        db_set = db_guids[guid_type]
        
        # GUIDs in XLSX but not in DB
        missing_in_db = xlsx_set - db_set
        
        # GUIDs in DB but not in XLSX
        missing_in_xlsx = db_set - xlsx_set
        
        results[guid_type] = {
            'xlsx_count': len(xlsx_set),
            'db_count': len(db_set),
            'missing_in_db': len(missing_in_db),
            'missing_in_xlsx': len(missing_in_xlsx),
            'invalid_guids': list(missing_in_db)[:100],  # First 100 examples
            'coverage': f"{(len(xlsx_set - missing_in_db) / len(xlsx_set) * 100):.1f}%" if xlsx_set else "0%"
        }
        
        print(f"   {guid_type:15} - XLSX: {len(xlsx_set):6,} | DB: {len(db_set):6,} | Missing: {len(missing_in_db):6,} | Coverage: {results[guid_type]['coverage']}")
    
    return results


def main():
    print("[START] RS.GE WB Items - GUID Validation")
    print("=" * 80)
    
    # Load XLSX GUIDs
    file_path = r"D:\next-postgres-starter\Prompts\RS.GE - WB_Items_IN.xlsx"
    xlsx_guids = load_xlsx_and_extract_guids(file_path)
    if not xlsx_guids:
        print("[ERROR] Failed to load XLSX")
        return
    
    # Fetch DB GUIDs
    db_guids = fetch_db_guids()
    if not db_guids:
        print("[ERROR] Failed to fetch DB GUIDs")
        return
    
    # Validate
    results = validate_guids(xlsx_guids, db_guids)
    
    # Print summary
    print(f"\n" + "=" * 80)
    print("VALIDATION SUMMARY")
    print("=" * 80)
    
    total_invalid = sum(r['missing_in_db'] for r in results.values())
    
    for guid_type, result in results.items():
        status = "✓" if result['missing_in_db'] == 0 else "✗"
        print(f"\n{status} {guid_type.upper()}")
        print(f"   XLSX unique GUIDs: {result['xlsx_count']:,}")
        print(f"   DB records: {result['db_count']:,}")
        print(f"   Invalid/Missing in DB: {result['missing_in_db']:,}")
        print(f"   Coverage: {result['coverage']}")
        if result['invalid_guids']:
            print(f"   Examples of invalid GUIDs:")
            for guid in result['invalid_guids'][:5]:
                print(f"     - {guid}")
    
    print(f"\n" + "=" * 80)
    print(f"TOTAL INVALID GUIDs: {total_invalid:,}")
    
    if total_invalid == 0:
        print("✓ ALL GUIDS ARE VALID - Ready for sync!")
    else:
        print("✗ INVALID GUIDS FOUND")
    
    # Save detailed report
    report = {
        'timestamp': datetime.now().isoformat(),
        'summary': {guid_type: {k: v for k, v in result.items() if k != 'invalid_guids'} for guid_type, result in results.items()},
        'invalid_guids': {guid_type: result['invalid_guids'] for guid_type, result in results.items()}
    }
    
    with open('guid_validation_report_final.json', 'w', encoding='utf-8') as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    
    print(f"\nReport saved: guid_validation_report_final.json")
    print("\n[COMPLETE]")


if __name__ == '__main__':
    main()
