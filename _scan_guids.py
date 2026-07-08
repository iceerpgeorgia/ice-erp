import openpyxl
from openpyxl import load_workbook
from collections import defaultdict

file_path = r"D:\next-postgres-starter\Prompts\RS.GE - WB_Items_IN.xlsx"
wb = load_workbook(file_path, data_only=True)
ws = wb.active

# Count populated GUID columns
guid_stats = {
    'კონტრაგენტის GUID': 0,
    'პროექტი GUID': 0,
    'საქონელი GUID': 0,
    'კოდის GUID': 0
}

guid_columns = {
    39: 'კონტრაგენტის GUID',
    40: 'პროექტი GUID',
    41: 'საქონელი GUID',
    42: 'კოდის GUID'
}

total_rows = 0
rows_with_any_guid = 0

print("Scanning all rows for populated GUIDs...")
for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if row_idx == 1:  # Skip header
        continue
    
    total_rows += 1
    
    # Check if any GUID column has a value
    has_guid = False
    for col_idx, guid_name in guid_columns.items():
        if col_idx < len(row):
            value = row[col_idx]
            if value and str(value).strip():
                guid_stats[guid_name] += 1
                has_guid = True
    
    if has_guid:
        rows_with_any_guid += 1
        # Print first few rows with GUIDs
        if rows_with_any_guid <= 5:
            print(f"\n[Row {row_idx}] Has GUIDs:")
            for col_idx, guid_name in guid_columns.items():
                if col_idx < len(row):
                    value = row[col_idx]
                    if value:
                        print(f"  {guid_name}: {value}")

print(f"\n" + "="*80)
print(f"SUMMARY: Total data rows: {total_rows}")
print(f"Rows with at least one populated GUID: {rows_with_any_guid}")
print(f"\nPopulated GUID columns:")
for guid_name, count in guid_stats.items():
    pct = (count / total_rows * 100) if total_rows > 0 else 0
    print(f"  {guid_name}: {count:,} / {total_rows} ({pct:.1f}%)")

# Check coverage
all_populated = all(count == total_rows for count in guid_stats.values())
any_populated = any(count > 0 for count in guid_stats.values())

if all_populated:
    print("\n✓ ALL GUID columns are fully populated!")
elif any_populated:
    print("\n⚠ GUID columns are PARTIALLY populated (some rows have GUIDs, others don't)")
else:
    print("\n✗ NO GUID columns are populated (all empty)")
