"""
Extract unique project and financial code values from XLSX for analysis
"""
import openpyxl
from openpyxl import load_workbook
from collections import Counter
import json

file_path = 'D:\\next-postgres-starter\\Prompts\\RS.GE - WB_Items_IN.xlsx'
wb = load_workbook(file_path, data_only=True)
ws = wb.active

unique_projects = set()
unique_fc = set()
project_freq = Counter()
fc_freq = Counter()

row_count = 0
for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if row_idx == 1:
        continue
    
    row_count += 1
    
    # Project column (idx 1)
    if row[1]:
        proj_val = str(row[1]).strip()
        # Extract first part before |
        proj_name = proj_val.split('|')[0].strip() if '|' in proj_val else proj_val
        unique_projects.add(proj_name)
        project_freq[proj_name] += 1
    
    # Financial code column (idx 2)
    if row[2]:
        fc_val = str(row[2]).strip()
        # Extract code prefix
        import re
        match = re.match(r'^([0-9_\.]+)', fc_val)
        if match:
            fc_code = match.group(1)
        else:
            fc_code = fc_val
        unique_fc.add(fc_code)
        fc_freq[fc_code] += 1

print(f"Total rows processed: {row_count}")
print(f"\n=== UNIQUE PROJECTS ===")
print(f"Count: {len(unique_projects)}")
print("\nTop 20 projects by frequency:")
for proj, count in project_freq.most_common(20):
    print(f"  '{proj}': {count} occurrences")

print(f"\n=== UNIQUE FINANCIAL CODES ===")
print(f"Count: {len(unique_fc)}")
print("\nAll financial codes by frequency:")
for fc, count in sorted(fc_freq.items(), key=lambda x: -x[1]):
    print(f"  '{fc}': {count} occurrences")

# Save to file for analysis
with open('D:\\next-postgres-starter\\xlsx_unique_values.json', 'w', encoding='utf-8') as f:
    json.dump({
        'unique_projects': sorted(list(unique_projects)),
        'unique_financial_codes': sorted(list(unique_fc)),
        'project_frequency': dict(project_freq.most_common(50)),
        'financial_code_frequency': dict(sorted(fc_freq.items(), key=lambda x: -x[1]))
    }, f, indent=2, ensure_ascii=False)

print("\n[OK] Unique values saved to: D:\\next-postgres-starter\\xlsx_unique_values.json")
