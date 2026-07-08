"""
RS.GE WB_Items sync - Fast batch insert version.
"""

import openpyxl
from openpyxl import load_workbook
import json
import os
import urllib.parse
from datetime import datetime
from dotenv import load_dotenv
import psycopg2
from psycopg2.extras import execute_batch

load_dotenv()

def parse_database_url(url):
    try:
        url_decoded = urllib.parse.unquote(url)
        result = urllib.parse.urlparse(url_decoded)
        return {
            'host': result.hostname,
            'port': result.port or 5432,
            'user': result.username,
            'password': result.password,
            'database': result.path.lstrip('/')
        }
    except Exception as e:
        print(f"[ERROR] {e}")
        return None


def main():
    print("[START] RS.GE Waybills Fast Sync")
    print("=" * 80)
    
    file_path = r"D:\next-postgres-starter\Prompts\RS.GE - WB_Items_IN.xlsx"
    
    # Load file
    print(f"[FILE] Loading...")
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active
    
    # Get headers and column indices
    headers = [cell.value for cell in ws[1]]
    col_map = {header: idx for idx, header in enumerate(headers)}
    
    print(f"   Headers: {len(headers)} columns")
    
    # Connect
    print(f"[DB] Connecting...")
    db_config = parse_database_url(os.getenv('DATABASE_URL'))
    conn = psycopg2.connect(**db_config)
    cur = conn.cursor()
    
    # Fetch validation dictionaries
    print(f"[DB] Fetching validation dictionaries...")
    
    cur.execute('SELECT LOWER(counteragent_uuid::text) FROM counteragents WHERE is_active = TRUE')
    valid_ca = set(row[0] for row in cur.fetchall())
    
    cur.execute('SELECT LOWER(project_uuid::text) FROM projects')
    valid_proj = set(row[0] for row in cur.fetchall())
    
    cur.execute('SELECT LOWER(uuid::text) FROM inventories WHERE is_active = TRUE')
    valid_inv = set(row[0] for row in cur.fetchall())
    
    cur.execute('SELECT LOWER(uuid::text) FROM financial_codes WHERE is_active = TRUE')
    valid_fc = set(row[0] for row in cur.fetchall())
    
    print(f"   Counteragents: {len(valid_ca):,}")
    print(f"   Projects: {len(valid_proj):,}")
    print(f"   Inventories: {len(valid_inv):,}")
    print(f"   Codes: {len(valid_fc):,}")
    
    # Process rows
    print(f"\n[SYNC] Processing {ws.max_row - 1:,} rows...")
    
    stats = {
        'total': ws.max_row - 1,
        'processed': 0,
        'inserted': 0,
        'updated': 0,
        'skipped': 0
    }
    
    issues = {
        'invalid_ca': [],
        'invalid_proj': [],
        'invalid_inv': [],
        'invalid_fc': [],
        'missing_fields': []
    }
    
    # Batch for inserts
    batch_inserts = []
    
    for row_idx in range(2, ws.max_row + 1):
        if row_idx % 5000 == 0:
            print(f"   Row {row_idx}...")
        
        # Get values
        waybill_no = ws.cell(row=row_idx, column=col_map.get('ზედნადების ნომერი', 0) + 1).value
        goods_code = ws.cell(row=row_idx, column=col_map.get('საქონლის კოდი', 0) + 1).value
        goods_name = ws.cell(row=row_idx, column=col_map.get('საქონლის დასახელება', 0) + 1).value
        ca_guid = ws.cell(row=row_idx, column=col_map.get('კონტრაგენტის GUID', 0) + 1).value
        proj_guid = ws.cell(row=row_idx, column=col_map.get('პროექტი GUID', 0) + 1).value
        inv_guid = ws.cell(row=row_idx, column=col_map.get('საქონელი GUID', 0) + 1).value
        fc_guid = ws.cell(row=row_idx, column=col_map.get('კოდის GUID', 0) + 1).value
        
        # Normalize
        if ca_guid and isinstance(ca_guid, str):
            ca_guid = ca_guid.strip().lower()
        if proj_guid and isinstance(proj_guid, str):
            proj_guid = proj_guid.strip().lower()
        if inv_guid and isinstance(inv_guid, str):
            inv_guid = inv_guid.strip().lower()
        if fc_guid and isinstance(fc_guid, str):
            fc_guid = fc_guid.strip().lower()
        
        # Validate
        if ca_guid and ca_guid not in valid_ca:
            issues['invalid_ca'].append({'row': row_idx, 'guid': ca_guid, 'waybill': waybill_no})
            stats['skipped'] += 1
            continue
        
        if proj_guid and proj_guid not in valid_proj:
            issues['invalid_proj'].append({'row': row_idx, 'guid': proj_guid, 'waybill': waybill_no})
            stats['skipped'] += 1
            continue
        
        if inv_guid and inv_guid not in valid_inv:
            issues['invalid_inv'].append({'row': row_idx, 'guid': inv_guid, 'waybill': waybill_no})
            stats['skipped'] += 1
            continue
        
        if fc_guid and fc_guid not in valid_fc:
            issues['invalid_fc'].append({'row': row_idx, 'guid': fc_guid, 'waybill': waybill_no})
            stats['skipped'] += 1
            continue
        
        if not waybill_no or not goods_code:
            issues['missing_fields'].append({'row': row_idx, 'waybill': waybill_no, 'code': goods_code})
            stats['skipped'] += 1
            continue
        
        # Add to batch
        if inv_guid:
            goods_code_str = str(int(goods_code) if isinstance(goods_code, (int, float)) else goods_code)
            batch_inserts.append((
                str(waybill_no),
                goods_code_str,
                goods_name,
                inv_guid,
                fc_guid
            ))
        
        stats['processed'] += 1
    
    print(f"\n   Total processed: {stats['processed']:,}")
    print(f"   Total skipped: {stats['skipped']:,}")
    
    # Insert all at once
    print(f"\n[INSERT] Inserting {len(batch_inserts):,} rows...")
    
    try:
        # Use the insider_uuid that's already used in the table
        # (This is the default insider for system-generated records)
        default_insider_uuid = '2a55debb-261b-4ce9-bae4-296ddea037ab'
        
        # Add insider_uuid to batch inserts
        batch_with_insider = [(
            waybill, code, name, inv, fc, default_insider_uuid
        ) for waybill, code, name, inv, fc in batch_inserts]
        
        cur = conn.cursor()
        execute_batch(cur, '''
            INSERT INTO rs_waybills_in_items (
                waybill_no, goods_code, goods_name, inventory_uuid, financial_code_uuid, insider_uuid, created_at, updated_at
            ) VALUES (%s, %s, %s, %s, %s, %s, NOW(), NOW())
        ''', batch_with_insider, page_size=1000)
        
        conn.commit()
        stats['inserted'] = len(batch_inserts)
        print(f"   [OK] Inserted: {len(batch_inserts):,}")
    except Exception as e:
        print(f"   [ERROR] {e}")
        conn.rollback()
    finally:
        cur.close()
        conn.close()
    
    # Report
    print(f"\n" + "=" * 80)
    print("SYNC REPORT")
    print("=" * 80)
    print(f"\nSTATS:")
    print(f"  Total rows: {stats['total']:,}")
    print(f"  Processed: {stats['processed']:,}")
    print(f"  Inserted: {stats['inserted']:,}")
    print(f"  Skipped: {stats['skipped']:,}")
    
    total_issues = sum(len(v) for v in issues.values())
    print(f"\nINCONSISTENCIES: {total_issues}")
    
    if issues['invalid_ca']:
        print(f"\n  ✗ Invalid Counteragent GUIDs: {len(issues['invalid_ca'])}")
        for issue in issues['invalid_ca'][:3]:
            print(f"    Row {issue['row']}: {issue['guid'][:20]}... (Waybill: {issue['waybill']})")
    
    if issues['invalid_proj']:
        print(f"\n  ✗ Invalid Project GUIDs: {len(issues['invalid_proj'])}")
        for issue in issues['invalid_proj'][:3]:
            print(f"    Row {issue['row']}: {issue['guid']} (Waybill: {issue['waybill']})")
    
    if issues['invalid_inv']:
        print(f"\n  ✗ Invalid Inventory GUIDs: {len(issues['invalid_inv'])}")
        for issue in issues['invalid_inv'][:3]:
            print(f"    Row {issue['row']}: {issue['guid'][:20]}... (Waybill: {issue['waybill']})")
    
    if issues['invalid_fc']:
        print(f"\n  ✗ Invalid Code GUIDs: {len(issues['invalid_fc'])}")
        for issue in issues['invalid_fc'][:3]:
            print(f"    Row {issue['row']}: {issue['guid']} (Waybill: {issue['waybill']})")
    
    if issues['missing_fields']:
        print(f"\n  ⚠ Missing Required Fields: {len(issues['missing_fields'])}")
        for issue in issues['missing_fields'][:3]:
            print(f"    Row {issue['row']}: waybill={issue['waybill']}, code={issue['code']}")
    
    # Save report
    report = {
        'timestamp': datetime.now().isoformat(),
        'stats': stats,
        'issues': {
            'invalid_counteragent_guids': issues['invalid_ca'][:100],
            'invalid_project_guids': issues['invalid_proj'][:100],
            'invalid_inventory_guids': issues['invalid_inv'][:100],
            'invalid_code_guids': issues['invalid_fc'][:100],
            'missing_required_fields': issues['missing_fields'][:100]
        }
    }
    
    with open('waybill_sync_report.json', 'w', encoding='utf-8') as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    
    print(f"\n[OK] Report: waybill_sync_report.json")
    print("[COMPLETE]")


if __name__ == '__main__':
    main()
